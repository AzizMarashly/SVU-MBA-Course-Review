# -*- coding: utf-8 -*-
"""Stage 3 support: text of every non-book source file -> ../extracted/srcNN/text.txt (+ index.json).
PDF via PyMuPDF with the lam-alef fix; DOCX via python-docx (paragraphs + tables, embedded images listed);
TXT copied. Scanned PDFs (no text layer) get a note and must be read visually (pages rendered by the helper).
Usage (from render/): set PYTHONUTF8=1 && python extract_sources.py"""
import os, re, json, io, zipfile
import pymupdf
HERE = os.path.dirname(os.path.abspath(__file__)); WORK = os.path.dirname(HERE); COURSE = os.path.dirname(WORK)
def fix(t): return re.sub("ا([أإآا])ل", r"ال\1", t)
rows = json.load(open(os.path.join(WORK, "extracted", "hashes.json"), encoding="utf-8"))
index = []
n = 0
for r in rows:
    p = r["path"]
    if p.startswith("Dr Iyad") or p.startswith("ملفات بوربوينت"): continue
    n += 1; sid = f"src{n:02d}"; d = os.path.join(WORK, "extracted", sid); os.makedirs(d, exist_ok=True)
    full = os.path.join(COURSE, r["disk_path"]); ext = p.rsplit(".", 1)[-1].lower(); note = ""; pages = None; text = ""
    if ext == "pdf":
        doc = pymupdf.open(full); pages = doc.page_count
        parts = []; empty = 0
        for i, pg in enumerate(doc, 1):
            t = fix(pg.get_text())
            if len(t.strip()) < 20: empty += 1
            parts.append(f"\n=== PAGE {i} ===\n{t}")
        text = "".join(parts)
        if empty == pages: note = "scanned: no text layer, read visually"
        elif empty: note = f"{empty} of {pages} pages have no text layer"
    elif ext == "docx":
        import docx
        dd = docx.Document(full); parts = [pa.text for pa in dd.paragraphs]
        for tb in dd.tables:
            for row in tb.rows: parts.append(" | ".join(c.text for c in row.cells))
        text = "\n".join(parts)
        imgs = [nm for nm in zipfile.ZipFile(full).namelist() if nm.startswith("word/media/")]
        note = f"{len(imgs)} embedded images" if imgs else "no embedded images"
        for nm in imgs:
            with open(os.path.join(d, os.path.basename(nm)), "wb") as f: f.write(zipfile.ZipFile(full).read(nm))
    elif ext == "txt":
        text = open(full, encoding="utf-8", errors="replace").read()
    elif ext == "jpg":
        note = "photo, read visually"
    with open(os.path.join(d, "text.txt"), "w", encoding="utf-8", newline="\n") as f: f.write(text)
    index.append(dict(id=sid, path=p, md5=r["md5"], pages=pages, chars=len(text), note=note))
    print(sid, pages, len(text), note, p)
json.dump(index, open(os.path.join(WORK, "extracted", "index.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
